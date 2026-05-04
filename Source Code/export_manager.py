import os
from PyQt6.QtGui import QTextDocument, QPageLayout
from PyQt6.QtPrintSupport import QPrinter

class ExportManager:
    """
    Dedicated class for handling all professional document generation and exporting.
    Separates the UI logic from the heavy formatting logic.
    """
    def __init__(self, timetable_data, mode="section"):
        self.timetable_data = timetable_data
        self.mode = mode
        
        # Extract unique constraints for pagination and sheets
        self.sections = set()
        self.buildings = set()
        self.teachers = set()
        
        for c in self.timetable_data:
            for s in c.get('sections', []):
                self.sections.add(s)
            if c.get('building'):
                self.buildings.add(c['building'])
            if c.get('instructors'):
                for t in c['instructors']:
                    self.teachers.add(t)
            elif c.get('instructor'):
                self.teachers.add(c['instructor'])
                
        self.sections = sorted(list(self.sections))
        self.buildings = sorted(list(self.buildings))
        self.teachers = sorted(list(self.teachers))

    def export_pdf(self, file_path: str):
        """Generates a professionally formatted multi-page PDF (one grid per section)."""
        printer = QPrinter(QPrinter.PrinterMode.ScreenResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(file_path)
        
        layout = printer.pageLayout()
        layout.setOrientation(QPageLayout.Orientation.Landscape)
        printer.setPageLayout(layout)
        
        doc = QTextDocument()
        html = ""
            
        if self.mode == "section":
            for idx, group in enumerate(self.sections):
                html += f"<h2 style='text-align: center; color: #2C3E50; font-family: sans-serif; letter-spacing: 1px;'>University Timetable</h2>"
                html += f"<h3 style='text-align: center; color: #34495E; font-family: sans-serif;'>Program: {group}</h3><br>"
                
                # Draw HTML Grid
                html += "<table border='1' cellspacing='0' cellpadding='10' width='100%' style='border-collapse: collapse; text-align: center; font-family: Arial, sans-serif; font-size: 13px;'>"
                html += "<tr bgcolor='#ecf0f1'><th style='color:#2c3e50;'>Day</th><th>8:00-8:50</th><th>9:00-9:50</th><th>10:30-11:20</th><th>11:30-12:20</th><th>12:30-1:20</th><th>2:30-3:20</th><th>3:30-4:20</th><th>4:30-5:20</th></tr>"
                
                days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
                for day_name in days:
                    html += f"<tr><td bgcolor='#f8f9fa' style='color:#2c3e50;'><b>{day_name}</b></td>"
                    for s_idx in range(8):
                        cell_text = ""
                        for c in self.timetable_data:
                            if group in c.get('sections', []):
                                for slot in c['schedule']:
                                    if slot['day'] == day_name and slot['slot_index'] == s_idx:
                                        cell_text = f"<b>{c['course_code']}</b><br>{c['room']}<br><span style='font-size:11px; color:#555;'>({c['course_title']})</span>"
                        
                        bg_color = "#ffffff" if cell_text == "" else "#eef6fa"
                        html += f"<td bgcolor='{bg_color}'>{cell_text}</td>"
                    html += "</tr>"
                html += "</table>"
                
                # Page break for next section
                if idx < len(self.sections) - 1:
                    html += "<div style='page-break-after: always;'></div>"
                    
        elif self.mode == "building":
            for idx, bldg in enumerate(self.buildings):
                if not bldg: continue
                html += f"<h2 style='text-align: center; color: #2C3E50; font-family: sans-serif; letter-spacing: 1px;'>University Timetable</h2>"
                html += f"<h3 style='text-align: center; color: #34495E; font-family: sans-serif;'>Building: {bldg}</h3><br>"
                
                html += "<table border='1' cellspacing='0' cellpadding='6' width='100%' style='border-collapse: collapse; text-align: center; font-family: Arial, sans-serif; font-size: 11px;'>"
                html += "<tr bgcolor='#ecf0f1'><th style='color:#2c3e50;'>Day</th><th style='color:#2c3e50;'>Room</th><th>8:00-8:50</th><th>9:00-9:50</th><th>10:30-11:20</th><th>11:30-12:20</th><th>12:30-1:20</th><th>2:30-3:20</th><th>3:30-4:20</th><th>4:30-5:20</th></tr>"
                
                days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
                rooms_in_bldg = sorted(list({c['room'] for c in self.timetable_data if c.get('building') == bldg}))
                
                for day_name in days:
                    first_room = True
                    for room in rooms_in_bldg:
                        html += "<tr>"
                        if first_room:
                            html += f"<td rowspan='{len(rooms_in_bldg)}' bgcolor='#f8f9fa' style='color:#2c3e50; font-size:14px;'><b>{day_name}</b></td>"
                            first_room = False
                            
                        html += f"<td bgcolor='#f8f9fa' style='color:#2c3e50;'><b>{room}</b></td>"
                        
                        for s_idx in range(8):
                            cell_text = ""
                            for c in self.timetable_data:
                                if c['room'] == room and c.get('building') == bldg:
                                    for slot in c['schedule']:
                                        if slot['day'] == day_name and slot['slot_index'] == s_idx:
                                            cell_text = f"<b>{c['course_code']}</b><br><span style='color:#555;'>({', '.join(c.get('sections', []))})</span>"
                                            
                            bg_color = "#ffffff" if cell_text == "" else "#eef6fa"
                            html += f"<td bgcolor='{bg_color}'>{cell_text}</td>"
                        html += "</tr>"
                html += "</table>"
                
                if idx < len(self.buildings) - 1:
                    html += "<div style='page-break-after: always;'></div>"

        elif self.mode == "teacher":
            for idx, teacher in enumerate(self.teachers):
                if not teacher: continue
                html += f"<h2 style='text-align: center; color: #2C3E50; font-family: sans-serif; letter-spacing: 1px;'>University Timetable</h2>"
                html += f"<h3 style='text-align: center; color: #34495E; font-family: sans-serif;'>Teacher: {teacher}</h3><br>"
                
                # Draw HTML Grid
                html += "<table border='1' cellspacing='0' cellpadding='10' width='100%' style='border-collapse: collapse; text-align: center; font-family: Arial, sans-serif; font-size: 13px;'>"
                html += "<tr bgcolor='#ecf0f1'><th style='color:#2c3e50;'>Day</th><th>8:00-8:50</th><th>9:00-9:50</th><th>10:30-11:20</th><th>11:30-12:20</th><th>12:30-1:20</th><th>2:30-3:20</th><th>3:30-4:20</th><th>4:30-5:20</th></tr>"
                
                days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
                for day_name in days:
                    html += f"<tr><td bgcolor='#f8f9fa' style='color:#2c3e50;'><b>{day_name}</b></td>"
                    for s_idx in range(8):
                        cell_text = ""
                        for c in self.timetable_data:
                            inst_list = c.get('instructors', [c.get('instructor')])
                            if teacher in inst_list:
                                for slot in c['schedule']:
                                    if slot['day'] == day_name and slot['slot_index'] == s_idx:
                                        cell_text = f"<b>{c['course_code']}</b><br>{c.get('room', '')}<br><span style='font-size:11px; color:#555;'>({', '.join(c.get('sections', []))})</span>"
                        
                        bg_color = "#ffffff" if cell_text == "" else "#eef6fa"
                        html += f"<td bgcolor='{bg_color}'>{cell_text}</td>"
                    html += "</tr>"
                html += "</table>"
                
                # Page break for next section
                if idx < len(self.teachers) - 1:
                    html += "<div style='page-break-after: always;'></div>"
                
        doc.setHtml(html)
        doc.print(printer)

    def export_excel(self, file_path: str):
        """Generates a multi-sheet Excel file natively mapping lists, grids, and building constraints."""
        try:
            import openpyxl
            from openpyxl.styles import Alignment, PatternFill, Font, Border, Side
        except ImportError:
            import csv
            if file_path.endswith('.xlsx'): file_path = file_path[:-5] + '.csv'
            with open(file_path, 'w', newline='', encoding='utf-8') as f:
                writer = csv.writer(f)
                headers = ["Code", "Title", "Instructor", "Sections", "Building", "Room", "Schedule"]
                writer.writerow(["MASTER TIMETABLE"])
                writer.writerow(headers)
                for c in self.timetable_data:
                    sched_str = ", ".join([f"{s['day']} {s['time']}" for s in c['schedule']])
                    writer.writerow([
                        c['course_code'], c['course_title'], c['instructor'], 
                        ", ".join(c.get('sections', [])), c.get('building', ''), c.get('room', ''), sched_str
                    ])
                    
                days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
                if self.mode == "section":
                    for group in self.sections:
                        writer.writerow([])
                        writer.writerow([f"SECTION: {group}"])
                        writer.writerow(["Day", "8:00-8:50", "9:00-9:50", "10:30-11:20", "11:30-12:20", "12:30-1:20", "2:30-3:20", "3:30-4:20", "4:30-5:20"])
                        for day_name in days:
                            row_data = [day_name] + [""] * 8
                            for s_idx in range(8):
                                for c in self.timetable_data:
                                    if group in c.get('sections', []):
                                        for slot in c['schedule']:
                                            if slot['day'] == day_name and slot['slot_index'] == s_idx:
                                                row_data[s_idx + 1] = f"{c['course_code']} {c.get('room', '')} ({c.get('instructor', '')})"
                            writer.writerow(row_data)
                elif self.mode == "teacher":
                    for teacher in self.teachers:
                        writer.writerow([])
                        writer.writerow([f"TEACHER: {teacher}"])
                        writer.writerow(["Day", "8:00-8:50", "9:00-9:50", "10:30-11:20", "11:30-12:20", "12:30-1:20", "2:30-3:20", "3:30-4:20", "4:30-5:20"])
                        for day_name in days:
                            row_data = [day_name] + [""] * 8
                            for s_idx in range(8):
                                for c in self.timetable_data:
                                    inst_list = c.get('instructors', [c.get('instructor')])
                                    if teacher in inst_list:
                                        for slot in c['schedule']:
                                            if slot['day'] == day_name and slot['slot_index'] == s_idx:
                                                row_data[s_idx + 1] = f"{c['course_code']} {c.get('room', '')} ({', '.join(c.get('sections', []))})"
                            writer.writerow(row_data)
            return
            
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove default empty sheet
        
        # Styles
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                             top=Side(style='thin'), bottom=Side(style='thin'))
        header_fill = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
        cell_fill = PatternFill(start_color="EBF5FB", end_color="EBF5FB", fill_type="solid")
        header_font = Font(bold=True, color="2C3E50")
        
        # ==========================================
        # 1. MASTER TIMETABLE SHEET (List Format)
        # ==========================================
        ws_master = wb.create_sheet(title="Master Timetable")
        headers = ["Code", "Title", "Instructor", "Sections", "Building", "Room", "Schedule"]
        ws_master.append(headers)
        
        for c in self.timetable_data:
            # Condense schedule list into string
            sched_str = ", ".join([f"{s['day']} {s['time']}" for s in c['schedule']])
            ws_master.append([
                c['course_code'], c['course_title'], c['instructor'], 
                ", ".join(c['sections']), c['building'], c['room'], sched_str
            ])
            
        for cell in ws_master[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = thin_border
            
        for col in "ABCDEFG":
            ws_master.column_dimensions[col].width = 25
            
        # ==========================================
        # 2. CONDITIONAL SCHEDULES (Grid Format)
        # ==========================================
        days = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
        
        if self.mode == "section":
            for group in self.sections:
                safe_title = f"Sec_{group[:25]}".replace('/', '-').replace('\\', '-').replace('?', '').replace('*', '').replace('[', '').replace(']', '').replace(':', '-')
                ws = wb.create_sheet(title=safe_title)
                
                ws.append(["Day", "8:00-8:50", "9:00-9:50", "10:30-11:20", "11:30-12:20", "12:30-1:20", "2:30-3:20", "3:30-4:20", "4:30-5:20"])
                
                for day_name in days:
                    row_data = [day_name] + [""] * 8
                    for s_idx in range(8):
                        for c in self.timetable_data:
                            if group in c.get('sections', []):
                                for slot in c['schedule']:
                                    if slot['day'] == day_name and slot['slot_index'] == s_idx:
                                        row_data[s_idx + 1] = f"{c['course_code']}\n{c['room']}\n({c['instructor']})"
                    ws.append(row_data)
                    
                # Formatting Grid
                for row in ws.iter_rows(min_row=1, max_row=6, min_col=1, max_col=9):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
                        cell.border = thin_border
                        if cell.row == 1 or cell.column == 1:
                            cell.font = header_font
                            cell.fill = header_fill
                        elif cell.value:
                            cell.fill = cell_fill
                            
                ws.column_dimensions['A'].width = 15
                for col in "BCDEFGHI":
                    ws.column_dimensions[col].width = 18
                for r in range(2, 7):
                    ws.row_dimensions[r].height = 60
                    
        elif self.mode == "building":
            for bldg in self.buildings:
                if not bldg: continue
                safe_title = f"Bldg_{bldg[:25]}".replace('/', '-').replace('\\', '-').replace('?', '').replace('*', '').replace('[', '').replace(']', '').replace(':', '-')
                ws_bldg = wb.create_sheet(title=safe_title)
                
                # Identify unique rooms within this building
                rooms_in_bldg = sorted(list({c['room'] for c in self.timetable_data if c.get('building') == bldg}))
                
                ws_bldg.append(["Room", "Day", "8:00-8:50", "9:00-9:50", "10:30-11:20", "11:30-12:20", "12:30-1:20", "2:30-3:20", "3:30-4:20", "4:30-5:20"])
                
                current_row = 2
                for room in rooms_in_bldg:
                    for day_name in days:
                        row_data = [room, day_name] + [""] * 8
                        for s_idx in range(8):
                            for c in self.timetable_data:
                                if c['room'] == room and c.get('building') == bldg:
                                    for slot in c['schedule']:
                                        if slot['day'] == day_name and slot['slot_index'] == s_idx:
                                            row_data[s_idx + 2] = f"{c['course_code']}\n({', '.join(c.get('sections', []))})"
                        ws_bldg.append(row_data)
                        
                        # Row-level formatting
                        for col_idx, cell in enumerate(ws_bldg[current_row], 1):
                            cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
                            cell.border = thin_border
                            if col_idx <= 2 or current_row == 1:
                                cell.fill = header_fill
                                cell.font = header_font
                            elif cell.value:
                                cell.fill = cell_fill
                        current_row += 1
                
                ws_bldg.column_dimensions['A'].width = 15
                ws_bldg.column_dimensions['B'].width = 15
                for col in "CDEFGHIJ":
                    ws_bldg.column_dimensions[col].width = 18
                for r in range(2, current_row):
                    ws_bldg.row_dimensions[r].height = 45

        elif self.mode == "teacher":
            for teacher in self.teachers:
                if not teacher: continue
                safe_title = f"Tchr_{teacher[:25]}".replace('/', '-').replace('\\', '-').replace('?', '').replace('*', '').replace('[', '').replace(']', '').replace(':', '-')
                ws = wb.create_sheet(title=safe_title)
                
                ws.append(["Day", "8:00-8:50", "9:00-9:50", "10:30-11:20", "11:30-12:20", "12:30-1:20", "2:30-3:20", "3:30-4:20", "4:30-5:20"])
                
                for day_name in days:
                    row_data = [day_name] + [""] * 8
                    for s_idx in range(8):
                        for c in self.timetable_data:
                            inst_list = c.get('instructors', [c.get('instructor')])
                            if teacher in inst_list:
                                for slot in c['schedule']:
                                    if slot['day'] == day_name and slot['slot_index'] == s_idx:
                                        row_data[s_idx + 1] = f"{c['course_code']}\n{c.get('room', '')}\n({', '.join(c.get('sections', []))})"
                    ws.append(row_data)
                    
                # Formatting Grid
                for row in ws.iter_rows(min_row=1, max_row=6, min_col=1, max_col=9):
                    for cell in row:
                        cell.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
                        cell.border = thin_border
                        if cell.row == 1 or cell.column == 1:
                            cell.font = header_font
                            cell.fill = header_fill
                        elif cell.value:
                            cell.fill = cell_fill
                            
                ws.column_dimensions['A'].width = 15
                for col in "BCDEFGHI":
                    ws.column_dimensions[col].width = 18
                for r in range(2, 7):
                    ws.row_dimensions[r].height = 60

        # Finally save the massive compiled file
        wb.save(file_path)
